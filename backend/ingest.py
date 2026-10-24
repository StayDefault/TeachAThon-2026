"""Extract short, page-labelled source passages from uploaded study files."""

import io
import re
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pymupdf
from docx import Document
from PIL import Image
from pptx import Presentation

from ai import ocr_image


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_SOURCE_CHARS = 40_000
Image.MAX_IMAGE_PIXELS = 20_000_000


class IngestError(ValueError):
    pass


def extract(name, data):
    ext = Path(name).suffix.lower()
    if not data or len(data) > MAX_FILE_BYTES:
        raise IngestError(f"{name}: files must be between 1 byte and 10 MB")
    try:
        if ext in (".txt", ".md"):
            pages = [data.decode("utf-8-sig")]
            kind, label = "text", "Text notes"
        elif ext == ".pdf":
            doc = pymupdf.open(stream=data, filetype="pdf")
            pages = []
            for page in doc:
                text = page.get_text().strip()
                if not text:
                    area = page.rect.width * page.rect.height
                    if area <= 0:
                        raise IngestError(f"{name}: page has no readable area")
                    scale = min(1.5, (10_000_000 / area) ** 0.5)
                    text = ocr_image(page.get_pixmap(matrix=pymupdf.Matrix(scale, scale)).tobytes("png"), "image/png")
                pages.append(text)
            doc.close()
            kind, label = "pdf", "PDF"
        elif ext == ".docx":
            check_archive(data)
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
            pages = [text]
            kind, label = "doc", "Document"
        elif ext == ".pptx":
            check_archive(data)
            deck = Presentation(io.BytesIO(data))
            pages = ["\n".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
            kind, label = "doc", "Slides"
        elif ext in (".png", ".jpg", ".jpeg", ".webp"):
            image = Image.open(io.BytesIO(data))
            if image.width * image.height > 20_000_000:
                raise IngestError(f"{name}: image is too large")
            image.thumbnail((2000, 2000))
            output = io.BytesIO()
            image.convert("RGB").save(output, format="JPEG", quality=85)
            pages = [ocr_image(output.getvalue(), "image/jpeg")]
            kind, label = "image", "Image"
        else:
            raise IngestError(f"{name}: use PDF, image, DOCX, PPTX, TXT or MD")
    except (OSError, UnicodeError, ValueError, RuntimeError, BadZipFile) as exc:
        if isinstance(exc, IngestError):
            raise
        raise IngestError(f"{name}: could not read this file") from exc
    if not any(page.strip() for page in pages):
        raise IngestError(f"{name}: no readable text found")
    return {"kind": kind, "label": label, "pages": pages}


def check_archive(data):
    with ZipFile(io.BytesIO(data)) as archive:
        if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 50_000_000:
            raise IngestError("Document expands beyond the prototype's size limit")


def chunks_for(materials, extracted):
    chunks = []
    for material, result in zip(materials, extracted):
        for page_number, text in enumerate(result["pages"], 1):
            # shortcut: small files are scanned directly; add retrieval before accepting longer courses.
            pieces = re.split(r"(?<=[.!?])\s+|\n+", text)
            for piece in pieces:
                piece = piece.strip()
                for start in range(0, len(piece), 350):
                    quote = piece[start:start + 350].strip()
                    if quote:
                        chunks.append({"id": f"c{len(chunks) + 1}", "materialId": material["id"],
                                       "fileName": material["name"],
                                       "page": page_number if material["name"].lower().endswith((".pdf", ".pptx")) else None,
                                       "text": quote})
    if sum(len(c["text"]) for c in chunks) > MAX_SOURCE_CHARS:
        raise IngestError("The extracted material is too long; use a shorter lesson")
    return chunks
