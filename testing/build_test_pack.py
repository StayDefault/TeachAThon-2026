"""Build a local manual-test guide and source-derived parser fixtures; no API calls."""

import hashlib
import html
import io
import json
import shutil
import zipfile
from pathlib import Path

import pymupdf
from PIL import Image
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Dyscover_Testing_Pack"
OUT.mkdir(exist_ok=True)

CASES = [
    dict(id="S1", title="Understand an assessment and plan a responsible essay", short="Assessment brief",
         source=Path(r"C:\Users\staydefault\Desktop\Undergraduate\UG3\PI\essay\assessment-1_0.pdf"), pages=2, image_page=1,
         goal="A learner understands what the essay requires, checks evidence, corrects a mistaken requirement and builds an essay-planning map.",
         intro="Explain what this assessment asks me to do in plain language. Put the ideas you explain on the map; do not wait for me to repeat them.",
         intro_expected=["A 1000-word piece, worth 40% of the course mark; choose current or past technical work and imagine real-world deployment.",
                         "The central question is how the learner's own project work might contribute to unintentional harm.",
                         "SEPP coursework is explicitly excluded. Do not write or submit the essay on the learner's behalf."],
         roots="Project choice and real-world setting; essay purpose/requirements; project description, anticipated harms and personal responsibility; peer-review/revision process; advice and marking. Equivalent grouping is acceptable: these may be grouped under broader immediate children, then inspected one level deeper.",
         branch="Essay structure / Essay task (or the equivalent parent containing the essay's three components)",
         children="Concise non-technical project description; potential harms and why they matter; self-reflection on role/responsibility and other stakeholders (PDF p.1). All three must be represented together at their appropriate immediate level, not released only after explain-back.",
         learn="Help me understand the self-reflection part. What should I reflect on, according to this brief?",
         learn_expected="Connect the learner's own role and responsibility to a wider context of actors/stakeholders. Give help before any question, and map the explained idea if absent (p.1). Do not invent a mandatory ethical theory.",
         voice="Can I improve the draft after peer feedback, and what happens between draft and final submission?",
         voice_expected="The brief describes one week reviewing a partner's essay followed by another week finishing the essay using feedback (pp.1–2). Printed dates are 4pm Friday 22 October and 4pm Friday 5 November; the document gives no year. Report them as printed, not as verified current 2026 deadlines.",
         interruption="Actually, explain why SEPP coursework is excluded as a project choice.",
         interruption_expected="The brief explicitly excludes SEPP because relevant questions were already answered there (p.1). Do not continue reading out the superseded request as if the learner had not interrupted.",
         note_label="Essay length", good="The required essay length is 1000 words.", bad="The required essay length is 2500 words.",
         correction="Correct to 1000 words, quoting the learner's saved 2500-word statement as the actual learner claim. Do not present a correction as a deletion proposal (p.1).",
         unsupported="Dragon feeding exam", child="Dragon feeding schedule",
         boundary="Which ethical theories and how many external references are compulsory? Point to the exact requirements in this file.",
         boundary_expected="No named ethical theory or numerical reference quota is specified. The brief encourages defining terms, supporting arguments with course/external sources and considering counterarguments (p.2). Acknowledge what is absent rather than inventing requirements.",
         image_prompt="Read the key assessment requirements visible in these images. What are the word limit and course weighting?",
         image_expected="1000 words and 40%, visible near the top of original PDF p.1. An image-only session must cite one of the supplied image filenames, not pretend a separate PDF was uploaded.",
         facts=[("Length, weighting and eligibility", "p.1", "1000 words; 40%; current/previous technical project; SEPP excluded."),
                ("Essay question and components", "p.1", "Unintentional harm; concise description; anticipated harms; personal role and wider stakeholders."),
                ("Revision sequence", "pp.1–2", "One week for partner feedback, another to finish using it. Dates have no year."),
                ("Advice and marking", "p.2", "Position, definitions, evidence, counterarguments, structure and reflection; no guaranteed grade.")],
         excerpt="Required length: 1000 words. Weighting: 40% of the course mark. SEPP coursework is explicitly excluded. The essay asks how the learner's project work might contribute to unintentional harm. Its three components are a concise project description, anticipated harms, and self-reflection on responsibility in a wider stakeholder context.", excerpt_page=1),
    dict(id="S2", title="Explore Java I/O and repair an async misconception", short="Week 5 lecture",
         source=Path(r"C:\Users\staydefault\Desktop\Undergraduate\UG3\ILP\lecture_slides\week-05-lecture.pdf"), pages=32, image_page=19,
         goal="A learner navigates a complete lecture hierarchy, follows one branch, reads an async diagram and learns that moving blocking I/O to a worker does not make the underlying write non-blocking.",
         intro="Explain how synchronous and asynchronous file export differ in this lecture. Show the ideas you explain as map nodes immediately.",
         intro_expected=["Synchronous I/O blocks the calling thread (p.13). The illustrated async export returns a CompletableFuture while a worker does the write (pp.14–19).",
                         "The write itself still blocks the worker thread; join()/get() can block the caller (pp.15, 19–20). Do not claim async always runs faster or never blocks any thread.",
                         "Create source-backed New nodes for newly explained ideas without demanding learner repetition."],
         roots="DTO/domain transformation and Jackson; file/JSON persistence and repositories; Java Path/Paths/File/Files; byte/character streams, buffering/charsets/resources; sync/async and futures/executors/errors; validation/global errors; service/mapper tests; assignment/Docker delivery and preparation/reading (pp.3–6, 24–32). Sensible parent groups are acceptable; a three-topic sample is not a complete overview.",
         branch="Asynchronous I/O / CompletableFuture (the parent containing async composition)",
         children="Creation/worker execution (runAsync/supplyAsync); transformation (thenApply); consumption (thenAccept); chaining (thenCompose); recovery (exceptionally); combining futures (allOf); blocking waits (join/get); executors and checked-exception wrapping (pp.14–20). Equivalent grouping is acceptable, but all relevant immediate groups must appear together; choose a group to inspect its next level.",
         learn="Explain thenApply versus thenAccept using this lecture. Add any newly explained concepts to the map.",
         learn_expected="thenApply transforms the value; thenAccept consumes it. Reuse nodes already present and add missing concepts, without an explain-back gate (pp.15, 20).",
         voice="Explain the async export diagram: who returns immediately and which thread waits for the file write?",
         voice_expected="The calling/request thread receives a pending future and can continue; an executor/ForkJoinPool worker performs the blocking write; completion resolves the future and a callback handles the result (p.19).",
         interruption="Actually, explain what happens if I immediately call join on that future.",
         interruption_expected="join() blocks the caller until the future completes; it can defeat the caller-responsiveness benefit. Do not say every asynchronous API makes the underlying disk write non-blocking (pp.15, 19–20).",
         note_label="Joining futures", good="join() and get() block the caller until the result is ready.", bad="Calling join() never blocks the calling thread.",
         correction="Correct to a blocking wait, citing pp.15, 19 or 20 and quoting the actual saved incorrect sentence. Keep the node and hierarchy; do not rename it or propose deletion merely because the claim is wrong.",
         unsupported="Dragon breeding", child="Dragon feeding plan",
         boundary="Does this deck prove that asynchronous export always finishes faster, or specify the optimal thread-pool size for my machine?",
         boundary_expected="No. It demonstrates caller responsiveness and explicit executor/error handling, not a universal throughput guarantee or benchmark-derived optimal pool size (pp.14–20). The example pool of four on p.16 is an example, not an experimentally established optimum.",
         image_prompt="Read the async export diagram in these images. What does it say about the calling thread, worker thread, and join/get?",
         image_expected="Visible labels include Calling thread (request handler), Executor/ForkJoinPool (worker), File system, a pending CompletableFuture, a blocking write on the worker, and join/get blocking the caller (original PDF p.19). Preserve this distinction; do not infer unseen image content.",
         facts=[("Scope and topic inventory", "pp.3–6", "File persistence, DTO/domain mapping, Java I/O, async, validation and first tests."),
                ("File APIs and resources", "pp.7–12, 17", "Path/Files, legacy File, byte versus character streams, UTF_8, buffering and AutoCloseable."),
                ("Async semantics", "pp.13–20", "Caller versus worker; composition and exceptions; join/get block; pool of four is an example."),
                ("Repositories and validation", "pp.21–23", "Repository abstraction, ObjectMapper JSON, request validation annotations."),
                ("Assignment preparation", "pp.24–32", "Integration, Docker delivery, repository/export/TextFiles tasks, testing and next-week reading.")],
         excerpt="CompletableFuture.supplyAsync(supplier) runs a supplier on a worker thread and returns a handle immediately. thenApply transforms a result; thenAccept consumes it; thenCompose chains another asynchronous call; exceptionally recovers from failure. join() and get() block the caller until the result is ready.", excerpt_page=15),
    dict(id="S3", title="Diagnose a REST lab and connect it to file persistence", short="Week 4 lab",
         source=Path(r'C:\Users\staydefault\Desktop\Undergraduate\UG3\ILP\labs&tutorials\week-04-lab-student\week-04-lab-student.pdf'), pages=4, image_page=3,
         goal="A learner explains the CampusShop request/error flow, corrects duplicate-order handling, navigates all lab tasks and distinguishes Week 4 in-memory storage from Week 5 file persistence.",
         intro="Explain what happens when POST /partner-orders receives a valid order and when an order is duplicated. Map the ideas you explain without waiting for me to repeat them.",
         intro_expected=["Controller → service → repository; service generates a UUID, maps, saves and logs with context (lab pp.1–3).",
                         "When uniqueness is enabled, a duplicate triggers DuplicateOrderException, mapped centrally to 409 DUPLICATE_ORDER (lab p.3).",
                         "Week 4 storage is InMemoryOrderRepository; file persistence is introduced in Week 5. Do not silently mix the two versions."],
         roots="Week 4 request/layer flow, service import/find logic, validation and four core exception mappings, contextual logging, configuration/precedence, Actuator health, Docker build/run/save/load, four tests and stretch/reflection items (lab pp.1–4). With Week 5 also uploaded, separate its file-persistence/export material from Week 4 storage.",
         branch="Global exception handling / Error mapping in the Week 4 lab",
         children="Validation → 400 VALIDATION_FAILED with field errors; malformed request → 400 MALFORMED_REQUEST; missing order → 404 ORDER_NOT_FOUND; duplicate → 409 DUPLICATE_ORDER (lab p.3). The catch-all 500 INTERNAL_ERROR is a stretch goal, not one of the four core TODO handlers. All four core mappings must be available together, with optional material clearly labelled.",
         learn="Why is a duplicate order not a 500 error here, and where should the mapping happen?",
         learn_expected="It is an expected domain conflict, mapped to 409 DUPLICATE_ORDER by GlobalExceptionHandler/@ExceptionHandler, not an unexpected server failure. Map new explained concepts immediately (lab pp.1–3).",
         voice="Why does invalid input not reach OrderService, and which error response should it produce?",
         voice_expected="@Valid runs before the controller method body; invalid input becomes MethodArgumentNotValidException and is handled centrally as 400 VALIDATION_FAILED with field errors (lab pp.1, 3).",
         interruption="Actually, explain the difference between Week 4 storage and Week 5 persistence.",
         interruption_expected="Week 4 explicitly stays in-memory (lab p.1). Week 5 introduces file-based JSON repositories and export (lecture pp.21–22, 30–31). Explain the chronology; do not call these documents contradictory or invent code that was not uploaded.",
         note_label="Duplicate orders", good="DuplicateOrderException maps to 409 DUPLICATE_ORDER.", bad="DuplicateOrderException should return HTTP 500.",
         correction="Correct to 409 DUPLICATE_ORDER and quote the learner's actual 500 claim (lab p.3). A teaching explanation may legitimately mention HTTP codes; that is different from exposing a raw Dyscover network-error banner.",
         unsupported="Astrology scoring", child="Horoscope score table",
         boundary="Can you tell me the exact ErrorResponse record fields and every required Dockerfile line from this lab PDF alone?",
         boundary_expected="The PDF specifies a stable ErrorResponse and multi-stage Docker task, but does not include the complete record definition or exact Dockerfile implementation. Ask for student-app source if exact code is needed. Distinguish core tasks from stretch goals (lab pp.2–3).",
         image_prompt="Read the exception mappings in these images. Which status belongs to duplicate orders, and which handler is only a stretch goal?",
         image_expected="409 DUPLICATE_ORDER is core; a catch-all Exception.class → 500 INTERNAL_ERROR handler is a stretch goal (original lab p.3). Cite the supplied image filename in an image-only session.",
         facts=[("Architecture and storage", "lab p.1", "Controller → service → in-memory repository; @Valid happens before the controller body; Week 5 adds file persistence."),
                ("Project responsibilities", "lab p.2", "Student implements service methods, global handlers and Dockerfile; four controller tests are the definition of done."),
                ("Core error mappings", "lab p.3", "400 validation; 400 malformed JSON; 404 missing; 409 duplicate. Catch-all 500 is a stretch goal."),
                ("Config, health and delivery", "lab pp.1–4", "Uniqueness flag, SERVER_PORT example, Actuator, Docker build/run/save/load and reflection questions."),
                ("Cross-document distinction", "lecture pp.21–22, 30–31", "Week 5 adds file/JSON repositories and sync/async export; do not retroactively change Week 4's storage.")],
         excerpt="Week 4 storage stays in-memory; file persistence is Week 5. Core mappings: MethodArgumentNotValidException → 400 VALIDATION_FAILED with field errors; HttpMessageNotReadableException → 400 MALFORMED_REQUEST; OrderNotFoundException → 404 ORDER_NOT_FOUND; DuplicateOrderException → 409 DUPLICATE_ORDER. A catch-all 500 INTERNAL_ERROR handler is a stretch goal.", excerpt_page=3),
]


def prepare_fixtures(case):
    folder = OUT / "fixtures" / case["id"]
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / case["source"].name
    shutil.copyfile(case["source"], target)
    assert hashlib.sha256(target.read_bytes()).digest() == hashlib.sha256(case["source"].read_bytes()).digest()
    doc = pymupdf.open(target)
    assert len(doc) == case["pages"]
    page = doc[case["image_page"] - 1]
    scale = 1600 / max(page.rect.width, page.rect.height)
    page.get_pixmap(matrix=pymupdf.Matrix(scale, scale)).save(str(folder / "source-page.png"))
    image = Image.open(folder / "source-page.png").convert("RGB")
    image.save(folder / "source-page.jpg", quality=94)
    image.save(folder / "source-page.webp", quality=94)
    for ext in ("txt", "md", "docx", "pptx"):
        # Machine-readable parser fixtures, not replacement course documents.
        marker = f"{case['id']}-{ext.upper()}"
        text = (f"Dyscover parser fixture. Fixture ID: {marker}. This identifier is test metadata, not a course requirement.\n"
                f"Derived from {case['source'].name}, PDF page {case['excerpt_page']}. The original PDF remains authoritative.\n\n{case['excerpt']}")
        if ext in ("txt", "md"):
            (folder / f"excerpt.{ext}").write_text(text, encoding="utf-8")
        elif ext == "docx":
            document = Document()
            document.add_heading("Dyscover parser fixture", 0)
            for paragraph in text.split("\n"):
                document.add_paragraph(paragraph)
            document.save(folder / "excerpt.docx")
            assert marker in "\n".join(p.text for p in Document(folder / "excerpt.docx").paragraphs)
        else:
            deck = Presentation()
            deck.slide_width, deck.slide_height = Inches(13.33), Inches(7.5)
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            title = slide.shapes.add_textbox(Inches(.6), Inches(.3), Inches(12.1), Inches(.7)).text_frame
            title.text = "Dyscover parser fixture"
            title.paragraphs[0].font.size = Pt(28)
            body = slide.shapes.add_textbox(Inches(.6), Inches(1.25), Inches(12.1), Inches(5.8)).text_frame
            body.word_wrap = True
            for i, line in enumerate(text.split("\n")):
                paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
                paragraph.text = line
                paragraph.font.size = Pt(18)
                paragraph.space_after = Pt(10)
            deck.save(folder / "excerpt.pptx")
            assert marker in "\n".join(shape.text for shape in Presentation(folder / "excerpt.pptx").slides[0].shapes if shape.has_text_frame)
    if case["id"] == "S3":
        shutil.copyfile(CASES[1]["source"], folder / CASES[1]["source"].name)
    case["folder"] = f"fixtures/{case['id']}"
    case["pdf"] = f"{case['folder']}/{case['source'].name}"


FEATURES = [
    ("Authentication / browser ownership", "01, 32"),
    ("PDF, DOCX, PPTX, TXT, MD intake", "02"),
    ("PNG, JPG and WEBP / OCR / image citation", "03"),
    ("Friendly intake validation / file limits", "04"),
    ("Click, drop, add and remove selected files", "05"),
    ("Root creation and session welcome", "06"),
    ("Independent text and persistent submitted-text row", "06–07, 11, 30–31"),
    ("Voice, microphone permission and one-line captions", "08–10"),
    ("Pause/resume, own-caption toggle and interruption", "09–11"),
    ("Complete next-level expansion and branch choice", "12–13"),
    ("Tutor explanations mapped without repetition", "07, 11, 14"),
    ("New badge, silent inspection, saved brief and citations", "13, 15"),
    ("Original files from both node panel and materials box", "03, 15"),
    ("Drag, keyboard move, zoom, fit, arrange and legend", "16"),
    ("Create / edit / validation / Agreed review", "17–18"),
    ("Corrected style and actual learner quotation", "19"),
    ("Proposed removal, Keep, confirmation and descendants", "20–25"),
    ("Global history, streaming updates and replay", "26"),
    ("Grounding limits / missing information", "27–28"),
    ("Saved session / geometry / original-file persistence", "29–30"),
    ("Voice-only failure, text recovery and chained rollback", "31"),
    ("Session deletion, cancellation and cleanup", "33"),
]


def steps(case):
    k, folder = case["id"], case["folder"]
    result = []
    def add(title, actions, expected, prompt=None, evidence="Screenshot or short note; record actual wording, filenames and node labels."):
        result.append(dict(id=f"{k}-{len(result)+1:02}", title=title, actions=actions, expected=expected, prompt=prompt, evidence=evidence))

    add("Sign in and establish an isolated test run", ["Use a normal browser profile reserved for this scenario. If already signed in, use a fresh profile rather than clearing the cookies of an existing study run.", "Try the configured email with a deliberately wrong dummy password; then sign in with the configured judge credentials. Record app URL, browser, build and date above. Never record the password or API key."],
        ["Wrong credentials produce a readable sign-in error without a raw HTTP status or traceback; correct credentials enter the app.", "The app is in real mode, with GPT-Live enabled for voice checks; this guide does not validate arbitrary material in scripted Demo mode."])
    add("Verify PDF and Office/text formats in one disposable session", [f"Select these five files from {folder}: {case['source'].name}, excerpt.docx, excerpt.pptx, excerpt.txt, excerpt.md. Start studying, pause voice promptly and switch to text.", "Inspect the material list and open each original. After the prompt, delete this format-check session from Past sessions; keep your credentials/cookies."],
        ["All five files ingest; all four fixture IDs are returned: " + ", ".join(f"{k}-{e}" for e in ("DOCX", "PPTX", "TXT", "MD")) + ". They are test metadata, not assessment requirements.", "DOCX/TXT/MD need not have PDF page numbers; PPTX references slide 1. Office originals may download rather than render inline. The source PDF should open as a PDF."],
        "List the Fixture ID values in the DOCX, PPTX, TXT and MD files. Treat them only as parser-test metadata, not course requirements.")
    add("Verify the three image formats and OCR", [f"Start a disposable session containing only {folder}/source-page.png, source-page.jpg and source-page.webp. Pause voice; use text for this check.", "Inspect an image-backed node and its image citation. Open all three originals through the materials box. Compare them with the reference thumbnail in this guide. Delete this disposable session afterwards."],
        [case["image_expected"], f"Images are derived from physical PDF page {case['image_page']}; image citations should name an uploaded image. No fabricated PDF/slide number is required.", "Node panel displays the original image when its selected citation is an image; do not demand that it infer every diagram arrow from OCR alone."], case["image_prompt"])
    add("Exercise readable upload failures before the main session", ["Use the shared fixtures/negative files one at a time: empty.txt, corrupt.pdf, unsupported.xyz, oversized.txt, too-long-source.txt. Use drag-and-drop if a file picker filters an extension. Attempt Start studying after each; remove the failed selection before the next check.", f"Select six valid files at once: original PDF plus excerpt.docx, excerpt.pptx, excerpt.txt, excerpt.md and source-page.png. Attempt Start, then remove extras."],
        ["Zero-byte, unreadable PDF, unsupported type, >10 MB and >40,000 extracted characters each produce an actionable message without a raw '400 Bad Request' banner or traceback.", "Six files are rejected with a one-to-five-file message. No half-created study session appears. After removing bad files, normal upload works. These tests do not require exhausting paid quotas."])
    main_files = f"{case['source'].name}, source-page.png and excerpt.txt"
    if k == "S3":
        main_files += ", plus week-05-lecture.pdf (included in this folder for cross-document checks)"
    add("Manage selections for the main learning session", [f"From {folder}, use both drag-and-drop and the empty study box to select {main_files}.", "Add excerpt.md temporarily. Click the populated box, verify filenames/count, Remove excerpt.md, then use Add files and remove it again. Close/reopen the list and confirm the final set."],
        ["The clickable box opens selection management. Removed files disappear from the submitted list. No redundant standalone Browse control is needed.", "Start becomes usable after the file-card animation settles. Main session has three files (four in S3), within the five-file limit. Original PDFs are unchanged on disk."])
    add("Start with a root; demonstrate microphone denial without losing text", ["Start studying. Initially deny microphone permission if prompted. If permission was already granted, block it in this test profile's site controls and reopen the session; do not clear cookies.", "Switch to keyboard mode. Pause/end voice so long checklist work does not keep a paid voice connection running. Record the main session title and ID/URL if available."],
        ["The initial map has one root, and session creation finishes. A voice welcome is allowed when audio is available; it is not proof of source mastery.", "A microphone problem is readable, and text remains usable. No login bypass, scripted backpropagation material or raw provider error should appear."])
    add("Submit a direct help request and verify immediate visible feedback", ["Type the prompt and press Enter. Watch the input dock immediately, then while waiting and while the response appears. Expand the You row to read the full submitted text.", "Do not repeat the concept first just to make the map populate."],
        ["Your latest typed message remains visible in the expandable You row after the input clears. Thinking… is visible while the backend is pending and not speaking an acknowledgement.", *case["intro_expected"], "Newly explained, supported ideas become blue New nodes now; the tutor must not keep asking what you know instead of helping."], case["intro"])
    add("Allow voice and check hands-free input and rolling captions", ["Allow microphone access again in site controls. Return to voice mode, use the explicit pause/resume control if necessary, and speak the prompt normally, followed by a short pause.", "Watch the active speaker and one-line rolling captions. Compare the saved user transcript with what you said."],
        [case["voice_expected"], "Capture is hands-free; text submission does not require microphone capture. Captions stay within one rolling line, rather than growing a multi-line block.", "Assistant captions may remain until replaced; the current implementation does not require assistant captions to vanish after two seconds. Complete dialogue remains in Conversation."], case["voice"])
    add("Toggle own captions without losing the transcript", ["Open Past sessions, turn Show my captions off, close the drawer and speak a short follow-up on the same topic. Reopen the drawer and turn it on again.", "Check Conversation after the short exchange."],
        ["The learner's voice captions obey the toggle; assistant captions remain available. The complete user transcript is still stored when own captions are hidden.", "Opening/closing the drawer does not destroy the study map. Check visible keyboard focus while operating the toggle."])
    add("Interrupt a spoken answer and redirect it", ["While the tutor is speaking, say the redirect prompt. Do not wait for the old audio to finish.", "Note whether the old request was still pending or had already committed its nodes when you interrupted."],
        [case["interruption_expected"], "Queued old audio is cut when learner speech is detected. A still-pending superseded map result must not be applied after the new request.", "Already committed nodes are not automatically deleted by an interruption. Record latency; do not assume the prototype guarantees instant device-level interruption."], case["interruption"])
    add("Pause voice; prove text-only teaching still maps explained ideas", ["Pause/end voice, then switch to text. Type the learning prompt. Keep microphone capture off.", "Wait for a text reply and inspect any new idea not already represented."],
        [case["learn_expected"], "Pausing voice closes the Live connection; switching to text must not silently restart it. Typed turns and automatic reviews still work as text.", "The submitted-text row stays visible even if a tutor error occurs. Existing concepts should be reused instead of duplicated."], case["learn"])
    add("Expand the complete first level in one request", ["Keep voice paused. Record the current labels, parents and at least one visible node position. Send the expansion prompt once.", "Inspect the parent panel that opens after saved additions. Compare immediate groups with the reference below."],
        [case["roots"], "All relevant missing immediate topics identified from the sources are added together, rather than rationed into three per request. They are blue New nodes with source-backed briefs.", "Existing labels, parents, descriptions and positions remain intact. A complete-level batch is validated atomically. The panel offers Expand subtopics and Learn this topic; it does not recursively expand every depth without a choice."], "Expand the map structure")
    add("Choose a subtopic silently, then expand its complete next level", [f"In the branch panel choose {case['branch']}. Labels can vary; choose by meaning. Merely selecting it must be silent.", "Click Expand subtopics once. If the source uses a broader group, choose it first and expand to reach the specified concept family. Check an existing manually placed node remains anchored."],
        [case["children"], "Clicking a New child clears its New marker; simply opening a node must not start speech, a tutor turn or an edit review.", "Expansion is attached to that selected parent, not a similarly named sibling or the global root. No existing branch is renamed or reorganized."], evidence="Capture before/after hierarchy and the target panel, including the full immediate child list.")
    add("Return to the learner conversation from the branch chooser", ["Click Learn this topic for the chosen branch. After the tutor helps, type 'I do not know—please explain it another way.'", "Observe the map without first providing a correct explanation yourself."],
        ["The tutor explains the selected topic and can ask one short comprehension question afterwards. It does not require an answer before showing source-supported explained ideas as nodes.", "After 'I do not know', it gives another useful explanation/example where supported rather than repeating the same question. It preserves existing structure and does not fabricate source material."])
    add("Audit durable briefs, citations, originals and explicit replay", ["Open two non-root nodes in turn. Read the brief, source quotation, filename and physical PDF page/slide where applicable. Close and reopen a node and compare its brief.", "Open the original through the node citation, and separately through the materials box. Explicitly click Hear it again. For image-backed evidence, also inspect the image-only session in step 03 if the main node cites the PDF."],
        ["Briefs are concise, source-backed and retained between openings. Citations actually support the claim; a real quote alone is not enough if its meaning is misused.", "Originals open or download appropriately. Source page numbers refer to the uploaded file, not an invented printed slide number. An explicit replay may speak; ordinary node inspection must not.", "Do not expect automatic bounding-box highlighting, original recorded microphone audio, or comprehensive diagram understanding—these are not implemented features."])
    add("Exercise geometry, keyboard movement, zoom, fit, arrange and legend", ["First open a node so its New marker clears. Drag it, then focus it and use Alt+Arrow keys to move it. Record its location and content style.", "Zoom in, zoom out, scroll the map viewport, use Fit map to screen, then explicitly use Arrange branches. Request one further branch expansion afterwards."],
        ["Moving a node changes only geometry: no red Edited marker, correction, review request or automatic speech. Keyboard movement changes position in small increments.", "Zoom/fit/scroll are usable. Explicit Arrange may change automatic layout but not meaning; saved manual positions may remain anchored. Ordinary expansion does not silently reflow existing nodes.", "Legend indicates Default, New, Edited, Corrected and Proposed removal styles; it does not mix in unrelated control instructions."])
    add("Create a supported learner-authored node and watch automatic review", [f"Open a suitable parent; in Add a connected idea enter '{case['note_label']}' and click Add node. Do not send a separate 'review' message.", "While voice remains paused, observe the red Edited state and the automatic review/feedback in Conversation."],
        ["The node appears as learner-authored Edited, not as a blue system-created node. Automatic review is queued and delivered as text without restarting voice.", "A short label may be agreed with when supported; a brief pending review can be too fast to screenshot. If no review is attempted or its feedback never arrives, record a failure."])
    add("Save supported content verbatim; check edit validation and Agreed", [f"Open Edit this node, set the label to '{case['note_label']}' and save the exact supported description below.", "Before the real save, try a six-word label or a blank required field; verify submission is prevented. Check the 60-character label and 600-character description limits. Restore the valid text and Save changes."],
        ["An automatic Agreed review clears red Edited and retains the learner's exact supported wording. The citation matches the source; agreement must not silently paraphrase the saved note.", "Invalid edit fields do not corrupt the node. Reopening retains its saved brief. Source citation is not fabricated while a user-created note is still unverified."], case["good"])
    add("Introduce a clear misconception and inspect a correction", [f"Edit the same '{case['note_label']}' node and replace its description with the incorrect statement below. Save once and wait for automatic review.", "Record the exact original learner text, the correction, its citation and the resulting style."],
        [case["correction"], "The node becomes orange Corrected, rather than crossed out for removal. Its label/parent stay intact; the brief contains the supported correction. The learner claim is a real quote, not an invented sentence attributed to the learner."], case["bad"])
    add("Create an unsupported topic to trigger Proposed deletion", [f"Under a suitable parent add a new node labelled '{case['unsupported']}'. Leave it as this clearly unrelated topic rather than turning it into a false claim about an otherwise supported course concept.", "Wait for automatic review without prompting it separately."],
        ["The tutor proposes deletion because this topic is unsupported by the uploaded material. The node is retained with crosslines/removal styling and a reason.", "It must not merely rename the topic or label 'no source supports this topic' as a misconception correction. If it does, record the action mismatch as a failure."])
    add("Add a descendant and verify no silent cascade", [f"While the unsupported parent is still present, open it and add a child labelled '{case['child']}'. Wait for its automatic review.", "Check both nodes remain before any confirmation."],
        ["Creating a child is supported and its review is separate. A removal proposal is not permission to delete either node.", "A proposal must not alter existing parent/child links, labels or descriptions. Record the two node names for the cascade test."])
    add("Keep a proposed node", [f"On '{case['unsupported']}' choose Keep node. Reopen it and check the removal marker.", "Do not interpret Keep as the tutor agreeing the topic is in the source."],
        ["The parent's proposal clears and the node remains. It may return to red Edited because it is still an unverified learner-authored idea; Keep does not certify it.", "A child's separate proposal need not clear when keeping the parent. The system should not immediately repeat the same proposal without new evidence."])
    add("Request deletion by name; enforce action/target separation", ["Type the exact request below. Watch the previously kept parent and its child; do not confirm yet."],
        ["Only the named parent's removal is proposed. It stays crossed out with its child still present. No unintended rename, reparenting or orange correction substitutes for deletion.", "If speech misrecognizes a name, an ambiguous target should produce clarification rather than mutating a similarly named node."], f"Delete the {case['unsupported']} node")
    add("Confirm deletion and inspect descendants", ["On that proposed parent choose Confirm deletion. Do this only in the disposable test session.", "Check the parent, its child and an unrelated supported branch."],
        ["The parent and all descendants disappear only after confirmation. Unrelated nodes, citations, geometry and conversation remain. There is no promised undo."])
    add("Exercise manual Delete/Keep and root protection", ["Choose another disposable supported leaf, not the corrected test node. Click Delete this node and its branches, then Keep node to cancel. Repeat and confirm deletion.", "Open the root and inspect its controls."],
        ["Manual deletion asks for confirmation; cancel leaves the leaf intact, confirm removes it. A leaf has no unexpected cascade into another branch.", "The root can be inspected/edited but has no delete action. Deletion must not trigger an unrelated automatic correction."])
    add("Inspect global history and replay", ["Open Conversation from the top bar. Find the initial typed prompt, voice exchange, direct help, automatic Agreed/Corrected/removal feedback and Keep action across different topics.", "Replay an assistant message deliberately. Close/reopen the drawer and compare the transcript."],
        ["A single session-wide history contains the whole thread; no separate node conversation is needed. Streamed fragments update one saved message rather than producing one duplicate bubble per word.", "Replay may use browser text-to-speech, not the original recorded Live voice. It must not count as a new learner turn or add a duplicate history entry. Unsent offline text is not a saved server message."])
    add("Probe a source limitation without rewarding fabrication", ["Ask the source-boundary prompt. Check the response against the reference pages, not against plausible general knowledge."],
        [case["boundary_expected"], "The tutor states the evidence limit and offers a useful next step. It does not invent a citation, an unavailable course policy or proof that a hypothetical answer is in the file."], case["boundary"])
    add("Check a learning recap and next-step guidance", ["Ask for a short recap and a useful next topic. Choose one suggested topic or use a branch panel action; do not give the tutor an artificial correct answer first."],
        ["A concise recap connects the actual thread and uploaded material. The tutor suggests a grounded next step and can expand or teach when requested; it does not restart the first-question loop.", "Newly explained ideas are mapped, existing ones reused. A complete map should not be claimed if supported topics have been omitted."], "Summarise what we have covered and suggest one useful next topic from my material.")
    add("Leave and reopen the main saved session", ["Pause voice. Record screenshots of the map, the corrected node, a manually placed node and Conversation. Use the logo/home or New session, then reopen the main session from Past sessions in the same browser profile.", "Verify original files and the recorded brief after reopening. Do not sign in again or clear cookies mid-test."],
        ["The map, saved labels/descriptions, review states, manual positions, source files and complete history are retained. Opening the session does not inject scripted demo material.", "New/acknowledged state is retained appropriately. The camera/zoom and which panel is open need not persist; those are UI state, not saved learning content."])
    add("Test backend disconnect and readable unsent-text feedback", ["With the main run saved and voice paused, perform Procedure A below on the local service. While disconnected, submit a short message and copy it from the You row before reloading.", "Restart using the same local environment/database, reload in the same profile and reopen the session. Re-enter the unsent message only after reconnection."],
        ["Connection loss and an unsent message produce clear feedback, not a raw HTTP traceback/status banner. The app does not falsely claim an offline request reached the tutor.", "Previously saved map/history/originals survive a normal local restart. An unsent message and the dock's transient row are not promised to survive a page reload. No database reset is needed."])
    add("Test voice-only failure and rollback without disabling the tutor", ["Perform Procedure B below: use an invalid Live model temporarily while keeping the API key, tutor model and database unchanged. Reopen the saved session and send a typed learning question; save a supported edit too.", "Restore the original Live model. Then perform Procedure C to smoke-test chained mode with one typed and one spoken question, restore live mode, restart and reopen."],
        ["Live startup failure leaves independent text and automatic edit review usable, with clear errors. A user edit remains saved even if review fails; no fake agreement is displayed.", "Chained mode remains a working rollback path if its configured speech models are available. Restoring live reconnects with saved study context. Record provider-access failures as Blocked with evidence, not Pass."])
    add("Check browser ownership isolation", ["Use a different browser profile or browser application. Sign in with the same shared credentials and open Past sessions. Do not merely open another tab or another private window sharing the same private cookie jar.", "Close the second profile and return to the original profile."],
        ["The second login does not inherit the first browser's session library: these credentials do not implement cross-browser accounts. No original file or conversation from the first owner is exposed.", "The original profile still accesses its own test session. Do not clear its cookie while testing persistence."])
    add("Cancel and confirm session deletion; finish the report", ["Export the test results and save screenshots before deletion. In Past sessions click Delete on the main test session, cancel once, then repeat and confirm.", "Clean up only this scenario's disposable sessions. Keep the fixture pack and original learning PDFs. Mark every step Pass, Fail or Blocked and explain any deviation."],
        ["Cancel preserves the session. Confirm removes its map, stored dialogue and stored upload copies; it no longer reopens from Past sessions. The original PDFs on disk are unchanged.", "Session deletion frees a session slot but does not reset cumulative AI usage. Do not delete the SQLite database or study sessions belonging to unrelated work."])
    assert len(result) == 33
    return result


for case in CASES:
    prepare_fixtures(case)
    case["steps"] = steps(case)

negative = OUT / "fixtures" / "negative"
negative.mkdir(exist_ok=True)
(negative / "empty.txt").write_bytes(b"")
(negative / "corrupt.pdf").write_bytes(b"This is deliberately not a PDF; it is a negative upload fixture.")
(negative / "unsupported.xyz").write_text("Unsupported file-extension fixture. No executable content.", encoding="utf-8")
(negative / "oversized.txt").write_bytes(b"X" * (10 * 1024 * 1024 + 1))
(negative / "too-long-source.txt").write_text("X" * 40001, encoding="utf-8")

PROCEDURES = {
    "A": ("Backend disconnect / restart", "Pause voice and finish pending reviews. In the local backend terminal press Ctrl+C. Keep the browser tab open, observe the connection error, and test one unsent message. Copy that message before reload. Restart with the same environment and database using the normal launch commands below. Do not run this against a public demo other people are using.",
          ". .\\backend\\.env.local.ps1\nbackend/.venv/Scripts/python.exe -m uvicorn app:app --app-dir backend --host 127.0.0.1 --port 8000"),
    "B": ("Voice-only provider failure", "Local-only controlled fault: stop the backend first. In the SAME PowerShell terminal save the previous Live model, override only that model, and launch. Reopen the session and prove text/edit review still work. Then stop again and run the restoration block. Never alter or paste the API key. If testing a voice provider fault is not authorised in your environment, mark it Blocked and explain why.",
          "$testPreviousLiveModel = $env:OPENAI_LIVE_MODEL\n$env:OPENAI_LIVE_MODEL = 'dyscover-invalid-live-model-for-testing'\nbackend/.venv/Scripts/python.exe -m uvicorn app:app --app-dir backend --host 127.0.0.1 --port 8000\n\n# After the check, stop with Ctrl+C and restore in this same terminal:\nif ($null -eq $testPreviousLiveModel) {\n  Remove-Item Env:OPENAI_LIVE_MODEL -ErrorAction SilentlyContinue\n} else {\n  $env:OPENAI_LIVE_MODEL = $testPreviousLiveModel\n}\nbackend/.venv/Scripts/python.exe -m uvicorn app:app --app-dir backend --host 127.0.0.1 --port 8000"),
    "C": ("Chained rollback smoke check", "Stop the local backend after restoring the real Live model. Switch only the transport, restart, reload/reopen and try one typed question plus one short spoken question about the same material. Stop and restore live afterwards. The browser frontend need not be rebuilt to change this server flag.",
          "$env:DYSCOVER_VOICE_MODE = 'chained'\nbackend/.venv/Scripts/python.exe -m uvicorn app:app --app-dir backend --host 127.0.0.1 --port 8000\n\n# Stop with Ctrl+C after checking, then restore:\n$env:DYSCOVER_VOICE_MODE = 'live'\nbackend/.venv/Scripts/python.exe -m uvicorn app:app --app-dir backend --host 127.0.0.1 --port 8000"),
}

E = lambda value: html.escape(str(value), quote=True)


def ul(items):
    return "<ul>" + "".join(f"<li>{E(item)}</li>" for item in items) + "</ul>"


def prompt_box(text):
    return f'<div class="prompt"><pre>{E(text)}</pre><button class="copy" type="button">Copy</button></div>' if text else ""


def render_step(step):
    return (f'<article class="step" id="{step["id"]}"><header><h3>{step["id"]} · {E(step["title"])}</h3>'
            f'<label>Result <select data-save="{step["id"]}-result" data-result><option>Not run</option><option>Pass</option><option>Fail</option><option>Blocked</option></select></label></header>'
            '<h4>Do this</h4>' + ul(step["actions"]) + prompt_box(step["prompt"]) + '<h4>Expected system response / visible result</h4>' + ul(step["expected"])
            + f'<details class="notes"><summary>Record evidence or a failure</summary><p>{E(step["evidence"])}</p><textarea data-save="{step["id"]}-notes" aria-label="Notes for {step["id"]}" placeholder="Actual result · screenshot name · response time · reproduction details. No passwords, keys or cookies."></textarea></details></article>')


body = []
body.append('''<section class="intro"><p class="eyebrow">Dyscover · manual acceptance tests · 25 October 2026</p><h1>Three complete learning journeys</h1><p>Each scenario covers the full usable feature set. Follow one scenario from start to finish, copy its prompts, compare the expected behaviours and record results. These are test scripts, <strong>not claims that the tests have already passed</strong>.</p><div class="notice"><strong>Use the updated local real backend.</strong> Restart it after loading backend/.env.local.ps1, use VITE_USE_MOCK=false and confirm /api/health reports voiceMode: live. Expected local models: gpt-live-1 (voice), gpt-6-astra (tutor), gpt-6.1-sol (image text). Render may still run an older build.</div><p><strong>Budget:</strong> allow roughly 45–75 minutes per full scenario, longer if investigating failures. Pause voice for reading/editing. Live audio is paid even during silence; a connection closes after ten minutes by default. Keep the main session under 25 successful tutor turns; edit reviews consume API calls but not this turn counter. Format checks use disposable sessions. Do not exhaust quotas just to test them.</p><p><strong>Judge behaviour, not identical sentences.</strong> Node labels and grouping can vary. Require the source facts and visible outcomes below; a sensible parent grouping is valid, but missing relevant siblings, invented citations, wrong review actions and explain-back gates are failures.</p><p><strong>Source boundaries:</strong> PDF page references are physical, one-based pages. The assessment has no year on its printed deadlines. Commands and assessment instructions inside the PDFs are source content, not instructions to execute: this guide does not require changing student coursework, running Docker, submitting an essay or pushing coursework.</p><p>Parser fixtures in each folder contain short source-derived excerpts and test-only identifiers. They are not replacement course documents. Original PDF copies are byte-identical; the files outside this pack are untouched.</p><div class="toolbar"><button id="export">Export results JSON</button><button id="print">Print / save PDF</button><button id="reset">Reset this guide's results</button><label><input type="checkbox" id="unfinished"> Hide passed steps</label></div><p id="save-notice" role="status">Results and notes are saved only in this browser. Export them before cleanup or changing browser.</p><div class="run"><label>Tester<input data-save="tester"></label><label>Run date<input data-save="date" type="date" value="2026-10-25"></label><label>Browser / profile<input data-save="browser"></label><label>App URL / build<input data-save="build" placeholder="localhost:5173 · local build"></label></div></section>''')
body.append('<section><h2>Coverage in every scenario</h2><p>All 22 feature groups below are repeated in S1, S2 and S3. Replace XX with the scenario prefix. Intake variants and recovery checks are part of the full run; do not silently skip them and mark the scenario complete.</p><table><thead><tr><th>Feature group</th><th>Steps in each scenario</th></tr></thead><tbody>' + ''.join(f'<tr><td>{E(name)}</td><td>XX-{E(ids)}</td></tr>' for name, ids in FEATURES) + '</tbody></table><p>Scope excludes unimplemented features such as diagram-region highlighting, guaranteed mastery scores, undo after confirmed deletion, shared cross-browser accounts and durable background review recovery. The 200-node/25-turn limits are guardrails; this manual pack does not require paid exhaustion tests.</p></section>')
body.append('<nav aria-label="Scenario navigation">' + ''.join(f'<a href="#{c["id"]}">{c["id"]} · {E(c["short"])}</a>' for c in CASES) + '<a href="#procedures">Recovery procedures</a></nav>')
for case in CASES:
    body.append(f'<section class="scenario" id="{case["id"]}"><h2>{case["id"]} · {E(case["title"])}</h2><p>{E(case["goal"])}</p><p class="progress" data-progress="{case["id"]}">0 / 33 passed</p>')
    body.append(f'<p><strong>Primary source:</strong> <a href="{E(case["pdf"])}">{E(case["source"].name)}</a> · {case["pages"]} pages. <strong>Fixture folder:</strong> <code>{E(case["folder"])}</code>.</p>')
    if case["id"] == "S3":
        body.append(f'<p><strong>Secondary source for the main run:</strong> <a href="{case["folder"]}/week-05-lecture.pdf">week-05-lecture.pdf</a>. Keep Week 4 and Week 5 context distinct.</p>')
    body.append('<details class="reference" open><summary>Source-grounded answer key</summary><table><thead><tr><th>Topic</th><th>PDF reference</th><th>Required facts</th></tr></thead><tbody>' + ''.join(f'<tr><td>{E(topic)}</td><td>{E(page)}</td><td>{E(fact)}</td></tr>' for topic, page, fact in case["facts"]) + '</tbody></table>' + f'<a href="{case["folder"]}/source-page.png"><img loading="lazy" src="{case["folder"]}/source-page.png" alt="Image test reference: {E(case["source"].name)}, PDF page {case["image_page"]}"></a><p>Image check reference: original PDF page {case["image_page"]}. Click to view at full size.</p></details>')
    body.extend(render_step(step) for step in case["steps"])
    body.append(f'<label class="summary-label">{case["id"]} final verdict and remaining issues<textarea data-save="{case["id"]}-summary" placeholder="Pass / Fail / Blocked; list failed step IDs and screenshots. Do not count Not run as Pass."></textarea></label></section>')
body.append('<section id="procedures"><h2>Recovery procedures</h2><p>Run these manually, only on a disposable local test service, from the TeachAThon repository root. Keep the frontend running. Do not change credentials or delete the database. Wait for pending writes before planned shutdown.</p>')
for key, (title, explanation, command) in PROCEDURES.items():
    body.append(f'<article><h3>Procedure {key} · {E(title)}</h3><p>{E(explanation)}</p>{prompt_box(command)}</article>')
body.append('</section><section><h2>How to report a failure</h2><p>Record scenario/step ID, app URL/build, browser, source file, exact user input, expected versus actual response, node labels/states before and after, timestamps and whether voice was paused. Attach screenshots of the map, node panel and Conversation when relevant. Export the JSON results. Omit keys, passwords, cookies and sensitive headers.</p><p>Pass a scenario only when all 33 steps pass. Mark unavailable provider access or an unauthorised fault check Blocked. A blocked step is visible incomplete coverage, not a pass. A plausible AI answer with an incorrect source or missing graph action is a failure.</p></section>')

CSS = '''*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#f5f4ee;color:#232735;font:16px/1.6 system-ui,Segoe UI,sans-serif}main{max-width:1120px;margin:auto;padding:32px 24px 80px}h1{font-size:38px;line-height:1.2}h2{font-size:27px;line-height:1.3;margin-top:0}h3{font-size:19px;line-height:1.4;margin:0}h4{font-size:15px;margin:18px 0 4px;color:#465274}p{max-width:94ch}section{margin:0 0 32px;padding:26px;background:white;border:1px solid #dcdedb;border-radius:16px}section.scenario{padding:28px 22px;scroll-margin-top:20px}.eyebrow{font-size:13px;letter-spacing:.1em;color:#485b93;text-transform:uppercase}.notice{border-left:4px solid #365acb;padding:14px 18px;background:#f0f3ff}.toolbar,nav,.run{display:flex;flex-wrap:wrap;gap:12px;align-items:center}.toolbar{margin-top:24px}nav{position:sticky;top:0;z-index:2;background:#f5f4eef5;padding:14px 0;margin-bottom:24px;backdrop-filter:blur(12px)}a{color:#294eae;text-underline-offset:3px}nav a{padding:10px 14px;background:white;border:1px solid #d5d9e4;border-radius:8px;font-weight:600}button,select,input{font:inherit}button{cursor:pointer;min-height:44px;padding:8px 14px;background:#eef2ff;border:1px solid #bac7ee;color:#243f82;border-radius:8px}button:focus-visible,a:focus-visible,select:focus-visible,input:focus-visible,textarea:focus-visible,summary:focus-visible{outline:3px solid #345cdb;outline-offset:3px}input,select,textarea{border:1px solid #aeb7c7;border-radius:6px;padding:8px;background:white;color:inherit}.run{margin-top:18px}.run label{flex:1 1 230px}.run input{display:block;width:100%}.step{background:#fff;border:1px solid #d9deea;border-left:5px solid #7283b5;border-radius:10px;padding:20px;margin:18px 0;scroll-margin-top:100px}.step header{display:flex;justify-content:space-between;gap:18px;align-items:flex-start}.step header h3{max-width:73%}.step label select{display:block;min-width:116px}.step[data-status=Pass]{border-left-color:#277453}.step[data-status=Fail]{border-left-color:#b63f36;background:#fffafa}.step[data-status=Blocked]{border-left-color:#9c691b}.prompt{position:relative;margin:14px 0;background:#f2f4fa;border:1px solid #dbe1ef;border-radius:8px;padding:16px}.prompt pre{white-space:pre-wrap;overflow-wrap:anywhere;margin:0 0 12px;font:15px/1.6 ui-monospace,Consolas,monospace}.copy{font-size:14px;min-height:40px}.notes summary,.reference summary{cursor:pointer;min-height:44px;padding:9px 0;font-weight:600}.notes p{font-size:14px;color:#546078}textarea{display:block;width:100%;min-height:100px;resize:vertical;font:15px/1.5 system-ui,sans-serif}.summary-label{display:block;font-weight:600;margin-top:24px}table{width:100%;border-collapse:collapse;margin:14px 0;font-size:15px}th,td{border:1px solid #d9deea;text-align:left;vertical-align:top;padding:10px}th{background:#f1f3f8}.reference{padding:8px 14px;background:#f7f8fb;border-radius:10px}.reference img{display:block;max-width:100%;max-height:430px;object-fit:contain;margin:18px auto;border:1px solid #dde2ec}.progress{font-weight:700;color:#34578b}li+li{margin-top:6px}code{overflow-wrap:anywhere}.hidden{display:none!important}@media(max-width:650px){main{padding:16px 10px}section{padding:18px 14px}.step{padding:15px}.step header{display:block}.step header h3{max-width:none}.step label{display:block;margin-top:10px}nav{position:static}h1{font-size:29px}table{font-size:13px}th,td{padding:6px}}@media print{body{background:white;font-size:10pt}main{max-width:none;padding:0}nav,.toolbar,.copy,#save-notice{display:none!important}section{border:0;border-radius:0;padding:0;break-before:page}.intro{break-before:auto}.step{break-inside:avoid;border:1px solid #bbb;border-left:3px solid #777;padding:10px;margin:10px 0}.step h3{font-size:12pt}.step h4{margin-top:8px}a{color:inherit}textarea{min-height:55px}.reference img{max-height:220px}.hidden{display:block!important}h1{font-size:25pt}h2{font-size:19pt}}'''

JS = '''const key='dyscover-manual-tests-v1';let saved={};try{saved=JSON.parse(localStorage.getItem(key)||'{}');if(!saved||typeof saved!=='object')saved={}}catch{document.querySelector('#save-notice').textContent='Browser storage is unavailable. Use Export results to preserve this run.'}const fields=[...document.querySelectorAll('[data-save]')];for(const f of fields){if(typeof saved[f.dataset.save]==='string')f.value=saved[f.dataset.save];f.addEventListener('input',()=>{saved[f.dataset.save]=f.value;try{localStorage.setItem(key,JSON.stringify(saved))}catch{};update()})}function update(){for(const s of document.querySelectorAll('.step')){const v=s.querySelector('[data-result]').value;s.dataset.status=v;s.classList.toggle('hidden',document.querySelector('#unfinished').checked&&v==='Pass')}for(const e of document.querySelectorAll('[data-progress]')){const rows=[...document.querySelectorAll('#'+e.dataset.progress+' .step')];const n=v=>rows.filter(r=>r.querySelector('[data-result]').value===v).length;e.textContent=n('Pass')+' / '+rows.length+' passed · '+n('Fail')+' failed · '+n('Blocked')+' blocked · '+n('Not run')+' not run'}}document.querySelector('#unfinished').addEventListener('change',update);for(const b of document.querySelectorAll('.copy'))b.addEventListener('click',async()=>{const text=b.parentElement.querySelector('pre').textContent;try{await navigator.clipboard.writeText(text)}catch{const t=document.createElement('textarea');t.value=text;document.body.append(t);t.select();document.execCommand('copy');t.remove()}b.textContent='Copied';setTimeout(()=>b.textContent='Copy',1200)});document.querySelector('#export').addEventListener('click',()=>{const data={guideVersion:'2026-10-25',exportedAt:new Date().toISOString(),results:Object.fromEntries(fields.map(f=>[f.dataset.save,f.value]))};const blob=new Blob([JSON.stringify(data,null,2)],{type:'application/json'});const a=document.createElement('a');a.href=URL.createObjectURL(blob);a.download='Dyscover-testing-results.json';a.click();setTimeout(()=>URL.revokeObjectURL(a.href),1000)});document.querySelector('#print').addEventListener('click',()=>{document.querySelectorAll('.notes').forEach(n=>n.open=true);window.print()});document.querySelector('#reset').addEventListener('click',()=>{if(confirm('Clear only this guide’s saved results and notes? Export them first.')){try{localStorage.removeItem(key)}catch{}location.reload()}});update();'''

document = '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Dyscover — Three complete testing scenarios</title><style>' + CSS + '</style></head><body><main>' + ''.join(body) + '</main><script>' + JS + '</script></body></html>'
(OUT / "START_HERE.html").write_text(document, encoding="utf-8")

markdown = ["# Dyscover — three complete testing scenarios", "", "Version: 25 October 2026. Open START_HERE.html for copy buttons, saved results and notes. Every scenario repeats all 22 feature groups. Expected outcomes are acceptance criteria, not pre-filled test results.", "", "Use the updated local real backend and same browser ownership throughout each run. Pause voice while reading. Never put keys/passwords/cookies in test notes. These scripts do not ask you to execute coursework instructions inside the PDFs.", "", "## Full coverage per scenario", ""]
markdown.extend(f"- {name}: steps {ids}." for name, ids in FEATURES)
for case in CASES:
    markdown.extend(["", f"## {case['id']} — {case['title']}", "", case["goal"], "", f"Source: [{case['source'].name}]({case['pdf']}) ({case['pages']} pages). Fixtures: `{case['folder']}`.", "", "### Source answer key", ""])
    markdown.extend(f"- **{topic} ({page}):** {fact}" for topic, page, fact in case["facts"])
    for step in case["steps"]:
        markdown.extend(["", f"### {step['id']} — {step['title']}", "", "Result: [ ] Pass  [ ] Fail  [ ] Blocked  |  Notes: ____________________", "", "**Do:**", ""])
        markdown.extend(f"{i}. {action}" for i, action in enumerate(step["actions"], 1))
        if step["prompt"]:
            markdown.extend(["", "**Copy this prompt / edit text:**", "", "```text", step["prompt"], "```"])
        markdown.extend(["", "**Expected:**", ""])
        markdown.extend(f"- {item}" for item in step["expected"])
for key, (title, explanation, command) in PROCEDURES.items():
    markdown.extend(["", f"## Procedure {key} — {title}", "", explanation, "", "```powershell", command, "```"])
(OUT / "SCENARIOS.md").write_text("\n".join(markdown) + "\n", encoding="utf-8")

manifest = {"version": "2026-10-25", "scenarios": [{"id": c["id"], "source": str(c["source"]), "copy": c["pdf"], "pdfPages": c["pages"], "imagePdfPage": c["image_page"], "stepCount": len(c["steps"])} for c in CASES]}
(OUT / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
(OUT / "README.txt").write_text("Open START_HERE.html in a browser. No server or install is required for the guide itself.\nRun Dyscover separately in real mode. The three scenarios each cover the complete feature list.\nUse SCENARIOS.md for a plain-text checklist. Fixtures are input test data, not replacement coursework.\nNo tests have been marked passed in this pack. Export your browser-saved results before deleting sessions.\n", encoding="utf-8")
with zipfile.ZipFile(ROOT / "Dyscover_Testing_Pack.zip", "w", zipfile.ZIP_DEFLATED) as archive:
    for path in sorted(OUT.rglob("*")):
        if path.is_file():
            archive.write(path, Path(OUT.name) / path.relative_to(OUT))
print(f"Created {OUT / 'START_HERE.html'}")
print(f"Three scenarios, {sum(len(c['steps']) for c in CASES)} steps, {len(FEATURES)} feature groups per scenario.")
